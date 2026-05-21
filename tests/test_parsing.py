from __future__ import annotations

import io

import fitz
import pytest
from careerpilot.errors import AppError
from careerpilot.services.parsing import extract_text
from docx import Document


def test_extract_text_from_docx() -> None:
    document = Document()
    document.add_paragraph("Ada Lovelace")
    document.add_paragraph("Built agent workflow tooling")
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text("resume.docx", buffer.getvalue())

    assert "Ada Lovelace" in text
    assert "Built agent workflow tooling" in text


def test_extract_text_from_pdf() -> None:
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "CareerPilot PDF Resume")
    text = extract_text("resume.pdf", pdf.tobytes())

    assert "CareerPilot PDF Resume" in text


def test_extract_text_rejects_unknown_file_types() -> None:
    with pytest.raises(AppError) as exc:
        extract_text("resume.txt", b"plain text")

    assert exc.value.error_code == "UNSUPPORTED_FILE_TYPE"
